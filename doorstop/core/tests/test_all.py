# SPDX-License-Identifier: LGPL-3.0-only

"""Tests for the doorstop.core package."""

# pylint: disable=protected-access,unidiomatic-typecheck

import csv
import logging
import os
import pprint
import shutil
import tempfile
import unittest
from unittest.mock import Mock, patch

import openpyxl
import yaml

from doorstop import common, core
from doorstop.common import DoorstopError, DoorstopInfo, DoorstopWarning
from doorstop.core.builder import _clear_tree, _get_tree
from doorstop.core.tests import (
    EMPTY,
    ENV,
    FILES,
    FILES_MD,
    REASON,
    ROOT,
    SYS,
    DocumentNoSkip,
)
from doorstop.core.vcs import mockvcs


class TestItem(unittest.TestCase):
    """Integration tests for the Item class."""

    def setUp(self):
        self.path = os.path.join(FILES, "REQ001.yml")
        self.backup = common.read_text(self.path)
        self.item = core.Item(None, self.path)
        self.item.tree = Mock()
        self.item.tree.vcs = mockvcs.WorkingCopy(EMPTY)

    def tearDown(self):
        common.write_text(self.backup, self.path)

    def test_save_load(self):
        """Verify an item can be saved and loaded from a file."""
        self.item.level = "1.2.3"
        self.item.text = "Hello, world!"
        self.item.links = ["SYS001", "SYS002"]
        item2 = core.Item(None, os.path.join(FILES, "REQ001.yml"))
        self.assertEqual((1, 2, 3), item2.level)
        self.assertEqual("Hello, world!", item2.text)
        self.assertEqual(["SYS001", "SYS002"], item2.links)

    @unittest.skipUnless(os.getenv(ENV), REASON)
    def test_find_ref(self):
        """Verify an item's external reference can be found."""
        item = core.Item(None, os.path.join(FILES, "REQ003.yml"))
        item.tree = Mock()
        item.tree.vcs = mockvcs.WorkingCopy(FILES)
        path, line = item.find_ref()
        relpath = os.path.relpath(os.path.join(FILES, "external", "text.txt"), FILES)
        self.assertEqual(relpath, path)
        self.assertEqual(3, line)

    def test_find_ref_error(self):
        """Verify an error occurs when no external reference found."""
        self.item.ref = "not" "found"  # pylint: disable=implicit-str-concat
        self.assertRaises(DoorstopError, self.item.find_ref)

    @unittest.skipUnless(os.getenv(ENV), REASON)
    def test_find_references(self):
        """Verify an item's external reference can be found."""
        item = core.Item(None, os.path.join(FILES, "REQ006.yml"), ROOT)
        item.tree = Mock()
        item.tree.vcs = mockvcs.WorkingCopy(FILES)
        item.root = FILES
        ref_items = item.find_references()
        self.assertEqual(len(ref_items), 2)

        path1, keyword_line_1 = ref_items[0]
        relpath1 = os.path.relpath(os.path.join(FILES, "external", "text.txt"), FILES)
        self.assertEqual(path1, relpath1)
        self.assertEqual(keyword_line_1, 3)

        path2, keyword_line_2 = ref_items[1]
        relpath2 = os.path.relpath(os.path.join(FILES, "external", "text2.txt"), FILES)
        self.assertEqual(path2, relpath2)
        self.assertEqual(keyword_line_2, None)

    def test_find_ref_multiple_error(self):
        """Verify an error occurs when no external reference found."""
        self.item.references = [{"path": "not" "found"}]
        self.assertRaises(DoorstopError, self.item.find_references)


class TestItemMarkdown(unittest.TestCase):
    """Integration tests for the Item class storage Markdown format."""

    def setUp(self):
        self.path = os.path.join(FILES_MD, "REQ001.md")
        self.backup = common.read_text(self.path)
        self.item = core.Item(None, self.path, itemformat="markdown")
        self.item.tree = Mock()
        self.item.tree.vcs = mockvcs.WorkingCopy(EMPTY)

    def tearDown(self):
        common.write_text(self.backup, self.path)

    def test_header(self):
        """Verify header is parsed correctly"""
        item2 = core.Item(None, self.path, itemformat="markdown")
        self.assertEqual("Markdown Header", item2.header)

    def test_save_load(self):
        """Verify an item can be saved and loaded from a markdown file."""
        self.item.header = "Another Header"
        self.item.level = "1.2.3"
        self.item.text = "Hello, world!"
        self.item.links = ["SYS001", "SYS002"]
        item2 = core.Item(
            None, os.path.join(FILES_MD, "REQ001.md"), itemformat="markdown"
        )
        self.assertEqual("Another Header", item2.header)
        self.assertEqual((1, 2, 3), item2.level)
        self.assertEqual("Hello, world!", item2.text)
        self.assertEqual(["SYS001", "SYS002"], item2.links)


class TestDocument(unittest.TestCase):
    """Integration tests for the Document class."""

    def setUp(self):
        self.document = core.Document(FILES, root=ROOT)

    def tearDown(self):
        """Clean up temporary files."""
        for filename in os.listdir(EMPTY):
            path = os.path.join(EMPTY, filename)
            common.delete(path)

    def test_load(self):
        """Verify a document can be loaded from a directory."""
        doc = core.Document(FILES)
        self.assertEqual("REQ", doc.prefix)
        self.assertEqual("yaml", doc.itemformat)
        self.assertEqual(2, doc.digits)
        self.assertEqual(6, len(doc.items))

    def test_new(self):
        """Verify a new document can be created."""
        document = core.Document.new(None, EMPTY, FILES, prefix="SYS", digits=4)
        self.assertEqual("SYS", document.prefix)
        self.assertEqual("yaml", document.itemformat)
        self.assertEqual(4, document.digits)
        self.assertEqual(0, len(document.items))

    @patch("doorstop.settings.REORDER", False)
    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    def test_validate(self):
        """Verify a document can be validated."""
        self.assertTrue(self.document.validate())

    @patch("doorstop.settings.REORDER", False)
    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    def test_issues_count(self):
        """Verify a number of issues are found in a document."""
        issues = self.document.issues
        for issue in self.document.issues:
            logging.info(repr(issue))
        self.assertEqual(13, len(issues))

    @patch("doorstop.settings.REORDER", False)
    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    def test_issues_duplicate_level(self):
        """Verify duplicate item levels are detected."""
        expect = DoorstopWarning("duplicate level: 2.1 (REQ002, REQ2-001)")
        for issue in self.document.issues:
            logging.info(repr(issue))
            if type(issue) == type(expect) and issue.args == expect.args:
                break
        else:
            self.fail("issue not found: {}".format(expect))

    @patch("doorstop.settings.REORDER", False)
    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    def test_issues_skipped_level(self):
        """Verify skipped item levels are detected."""
        expect = DoorstopInfo("skipped level: 1.2.3 (REQ001), 1.4 (REQ003)")
        for issue in self.document.issues:
            logging.info(repr(issue))
            if type(issue) == type(expect) and issue.args == expect.args:
                break
        else:
            self.fail("issue not found: {}".format(expect))

    def test_add_item_with_reordering(self):
        """Verify an item can be inserted into a document."""
        document = core.Document.new(None, EMPTY, FILES, prefix="TMP")
        item_1_0 = document.add_item()
        item_1_2 = document.add_item()  # will get displaced
        item_1_1 = document.add_item(level="1.1")
        self.assertEqual((1, 0), item_1_0.level)
        self.assertEqual((1, 1), item_1_1.level)
        self.assertEqual((1, 2), item_1_2.level)

    def test_remove_item_with_reordering(self):
        """Verify an item can be removed from a document."""
        document = core.Document.new(None, EMPTY, FILES, prefix="TMP")
        item_1_0 = document.add_item()
        item_1_2 = document.add_item()  # to be removed
        item_1_1 = document.add_item(level="1.1")  # will get relocated
        document.remove_item(item_1_2)
        self.assertEqual((1, 0), item_1_0.level)
        self.assertEqual((1, 1), item_1_1.level)

    def test_reorder(self):
        """Verify a document's order can be corrected."""
        document = core.Document.new(None, EMPTY, FILES, prefix="TMP")
        document.add_item(level="2.0", reorder=False)
        document.add_item(level="2.1", reorder=False)
        document.add_item(level="2.1", reorder=False)
        document.add_item(level="2.5", reorder=False)
        document.add_item(level="4.5", reorder=False)
        document.add_item(level="4.7", reorder=False)
        document.reorder()
        expected = [(2, 0), (2, 1), (2, 2), (2, 3), (3, 1), (3, 2)]
        actual = [item.level for item in document.items]
        self.assertListEqual(expected, actual)

    def test_reorder_with_keep(self):
        """Verify a document's order can be corrected with a kept level."""
        document = core.Document.new(None, EMPTY, FILES, prefix="TMP")
        document.add_item(level="1.0", reorder=False)
        item = document.add_item(level="1.0", reorder=False)
        document.add_item(level="1.0", reorder=False)
        document.reorder(keep=item)
        expected = [(1, 0), (2, 0), (3, 0)]
        actual = [i.level for i in document.items]
        self.assertListEqual(expected, actual)
        self.assertEqual((1, 0), item.level)

    def test_reorder_with_start(self):
        """Verify a document's order can be corrected with a given start."""
        document = core.Document.new(None, EMPTY, FILES, prefix="TMP")
        document.add_item(level="2.0", reorder=False)
        document.add_item(level="2.1", reorder=False)
        document.add_item(level="2.1", reorder=False)
        document.add_item(level="2.5", reorder=False)
        document.add_item(level="4.0", reorder=False)
        document.add_item(level="4.7", reorder=False)
        document.reorder(start=(1, 0))
        expected = [(1, 0), (1, 1), (1, 2), (1, 3), (2, 0), (2, 1)]
        actual = [item.level for item in document.items]
        self.assertListEqual(expected, actual)

    @patch("doorstop.settings.REORDER", True)
    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    def test_validate_with_reordering(self):
        """Verify a document's order is corrected during validation."""
        document = core.Document.new(None, EMPTY, FILES, prefix="TMP")
        document.add_item(level="1.0", reorder=False)
        document.add_item(level="1.1", reorder=False)
        document.add_item(level="1.2.0", reorder=False)
        document.add_item(level="1.2.5", reorder=False)
        document.add_item(level="3.2.1", reorder=False)
        document.add_item(level="3.3", reorder=False)
        self.assertTrue(document.validate())
        expected = [(1, 0), (1, 1), (1, 2, 0), (1, 2, 1), (2, 1, 1), (2, 2)]
        actual = [item.level for item in document.items]
        self.assertListEqual(expected, actual)


class TestDocumentMarkdown(unittest.TestCase):
    """Integration tests for the Document class with Items in Markdown format."""

    def setUp(self):
        self.document = core.Document(FILES_MD, root=ROOT)

    def tearDown(self):
        """Clean up temporary files."""
        for filename in os.listdir(EMPTY):
            path = os.path.join(EMPTY, filename)
            common.delete(path)

    def test_load(self):
        """Verify a document can be loaded from a directory."""
        doc = core.Document(FILES_MD)
        self.assertEqual("REQMD", doc.prefix)
        self.assertEqual("markdown", doc.itemformat)
        self.assertEqual(2, doc.digits)
        self.assertEqual(1, len(doc.items))

    def test_new(self):
        """Verify a new document can be created."""
        document = core.Document.new(
            None, EMPTY, FILES_MD, prefix="SYSMD", digits=4, itemformat="markdown"
        )
        self.assertEqual("SYSMD", document.prefix)
        self.assertEqual("markdown", document.itemformat)
        self.assertEqual(4, document.digits)
        self.assertEqual(0, len(document.items))


class TestTree(unittest.TestCase):
    """Integration tests for the core.Tree class."""

    def setUp(self):
        self.path = os.path.join(FILES, "REQ001.yml")
        self.backup = common.read_text(self.path)
        self.item = core.Item(None, self.path)
        self.tree = core.Tree(core.Document(SYS))
        self.tree._place(core.Document(FILES))

    def tearDown(self):
        common.write_text(self.backup, self.path)

    @patch("doorstop.settings.REORDER", False)
    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    @patch("doorstop.settings.STAMP_NEW_LINKS", False)
    @patch("doorstop.settings.CHECK_REF", False)
    def test_issues_count(self):
        """Verify a number of issues are found in a tree."""
        issues = self.tree.issues
        for issue in self.tree.issues:
            logging.info(repr(issue))
        self.assertEqual(15, len(issues))

    @patch("doorstop.settings.REORDER", False)
    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    @patch("doorstop.settings.STAMP_NEW_LINKS", False)
    @patch("doorstop.settings.CHECK_REF", False)
    def test_issues_count_with_skips(self):
        """Verify a document can be skipped during validation."""
        issues = list(self.tree.get_issues(skip=["req"]))
        for issue in self.tree.issues:
            logging.info(repr(issue))
        self.assertEqual(2, len(issues))

    @patch("doorstop.settings.REORDER", False)
    @patch("doorstop.settings.STAMP_NEW_LINKS", False)
    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    @patch("doorstop.core.document.Document", DocumentNoSkip)
    def test_validate_invalid_link(self):
        """Verify a tree is invalid with a bad link."""
        self.item.link("SYS003")
        tree = core.build(FILES, root=FILES)
        self.assertIsInstance(tree, core.Tree)
        self.assertFalse(tree.validate())

    @patch("doorstop.settings.REORDER", False)
    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    def test_validate_long(self):
        """Verify trees can be checked."""
        logging.info("tree: {}".format(self.tree))
        self.assertTrue(self.tree.validate())


@unittest.skipUnless(os.getenv(ENV), REASON)
class TestEditor(unittest.TestCase):
    """Integrations tests for the editor module."""


class TestImporter(unittest.TestCase):
    """Integrations tests for the importer module."""

    def setUp(self):
        # Create a temporary mock working copy
        self.cwd = os.getcwd()
        self.temp = tempfile.mkdtemp()
        os.chdir(self.temp)
        common.touch(".mockvcs")
        # Create default document attributes
        self.prefix = "PREFIX"
        self.root = self.temp
        self.path = os.path.join(self.root, "DIRECTORY")
        self.parent = "PARENT_PREFIX"
        # Create default item attributes
        self.uid = "PREFIX-00042"
        # Load an actual document
        self.document = core.Document(FILES, root=ROOT)
        # Ensure the tree is reloaded
        _clear_tree()

    def tearDown(self):
        os.chdir(self.cwd)
        shutil.rmtree(self.temp)

    def test_import_yml(self):
        """Verify items can be imported from a YAML file."""
        path = os.path.join(self.temp, "exported.yml")
        core.exporter.export(self.document, path)
        _path = os.path.join(self.temp, "imports", "req")
        _tree = _get_tree()
        document = _tree.create_document(_path, "REQ")
        # Act
        core.importer.import_file(path, document)
        # Assert
        expected = [item.data for item in self.document.items]
        actual = [item.data for item in document.items]
        log_data(expected, actual)
        self.assertListEqual(expected, actual)

    def test_import_csv(self):
        """Verify items can be imported from a CSV file."""
        path = os.path.join(self.temp, "exported.csv")
        core.exporter.export(self.document, path)
        _path = os.path.join(self.temp, "imports", "req")
        _tree = _get_tree()
        document = _tree.create_document(_path, "REQ")
        # Act
        core.importer.import_file(path, document)
        # Assert
        expected = [item.data for item in self.document.items]
        actual = [item.data for item in document.items]
        log_data(expected, actual)
        self.assertListEqual(expected, actual)

    def test_import_tsv(self):
        """Verify items can be imported from a TSV file."""
        path = os.path.join(self.temp, "exported.tsv")
        core.exporter.export(self.document, path)
        _path = os.path.join(self.temp, "imports", "req")
        _tree = _get_tree()
        document = _tree.create_document(_path, "REQ")
        # Act
        core.importer.import_file(path, document)
        # Assert
        expected = [item.data for item in self.document.items]
        actual = [item.data for item in document.items]
        log_data(expected, actual)
        self.assertListEqual(expected, actual)

    @unittest.skipUnless(os.getenv(ENV), REASON)
    def test_import_xlsx(self):
        """Verify items can be imported from an XLSX file."""
        path = os.path.join(self.temp, "exported.xlsx")
        core.exporter.export(self.document, path)
        _path = os.path.join(self.temp, "imports", "req")
        _tree = _get_tree()
        document = _tree.create_document(_path, "REQ")
        # Act
        core.importer.import_file(path, document)
        # Assert
        expected = [item.data for item in self.document.items]
        actual = [item.data for item in document.items]
        log_data(expected, actual)
        self.assertListEqual(expected, actual)

    def test_create_document(self):
        """Verify a new document can be created to import items."""
        document = core.importer.create_document(self.prefix, self.path)
        self.assertEqual(self.prefix, document.prefix)
        self.assertEqual(self.path, document.path)

    def test_create_document_with_unknown_parent(self):
        """Verify a new document can be created with an unknown parent."""
        # Verify the document does not already exist
        self.assertRaises(DoorstopError, core.find_document, self.prefix)
        # Import a document
        document = core.importer.create_document(
            self.prefix, self.path, parent=self.parent
        )
        # Verify the imported document's attributes are correct
        self.assertEqual(self.prefix, document.prefix)
        self.assertEqual(self.path, document.path)
        self.assertEqual(self.parent, document.parent)
        # Verify the imported document can be found
        document2 = core.find_document(self.prefix)
        self.assertIs(document, document2)

    def test_create_document_already_exists(self):
        """Verify non-parent exceptions are re-raised."""
        # Create a document
        core.importer.create_document(self.prefix, self.path)
        # Attempt to create the same document
        self.assertRaises(
            DoorstopError, core.importer.create_document, self.prefix, self.path
        )

    def test_add_item(self):
        """Verify an item can be imported into a document."""
        # Create a document
        core.importer.create_document(self.prefix, self.path)
        # Verify the item does not already exist
        self.assertRaises(DoorstopError, core.find_item, self.uid)
        # Import an item
        item = core.importer.add_item(self.prefix, self.uid)
        # Verify the item's attributes are correct
        self.assertEqual(self.uid, item.uid)
        # Verify the item can be found
        item2 = core.find_item(self.uid)
        self.assertIs(item, item2)
        # Verify the item is contained in the document
        document = core.find_document(self.prefix)
        self.assertIn(item, document.items)

    def test_add_item_with_attrs(self):
        """Verify an item with attributes can be imported into a document."""
        # Create a document
        core.importer.create_document(self.prefix, self.path)
        # Import an item
        attrs = {"text": "Item text", "ext1": "Extended 1"}
        item = core.importer.add_item(self.prefix, self.uid, attrs=attrs)
        # Verify the item is correct
        self.assertEqual(self.uid, item.uid)
        self.assertEqual(attrs["text"], item.text)
        self.assertEqual(attrs["ext1"], item.get("ext1"))


class TestExporter(unittest.TestCase):
    """Integration tests for the doorstop.core.exporter module."""

    maxDiff = None

    def setUp(self):
        self.document = core.Document(FILES, root=ROOT)
        self.temp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.temp)

    def test_export_yml(self):
        """Verify a document can be exported as a YAML file."""
        path = os.path.join(FILES, "exported.yml")
        temp = os.path.join(self.temp, "exported.yml")
        expected = read_yml(path)
        # Act
        path2 = core.exporter.export(self.document, temp)
        # Assert
        self.assertIs(temp, path2)
        actual = read_yml(temp)
        self.assertEqual(expected, actual)
        move_file(temp, path)

    def test_export_csv(self):
        """Verify a document can be exported as a CSV file."""
        path = os.path.join(FILES, "exported.csv")
        temp = os.path.join(self.temp, "exported.csv")
        expected = read_csv(path)
        # Act
        path2 = core.exporter.export(self.document, temp)
        # Assert
        self.assertIs(temp, path2)
        actual = read_csv(temp)
        self.assertEqual(expected, actual)
        move_file(temp, path)

    @patch("doorstop.settings.REVIEW_NEW_ITEMS", False)
    def test_export_tsv(self):
        """Verify a document can be exported as a TSV file."""
        path = os.path.join(FILES, "exported.tsv")
        temp = os.path.join(self.temp, "exported.tsv")
        expected = read_csv(path, delimiter="\t")
        # Act
        path2 = core.exporter.export(self.document, temp)
        # Assert
        self.assertIs(temp, path2)
        actual = read_csv(temp, delimiter="\t")
        self.assertEqual(expected, actual)
        move_file(temp, path)


class TestPublisher(unittest.TestCase):
    """Integration tests for the doorstop.core.publisher module."""

    maxDiff = None

    @patch("doorstop.core.document.Document", DocumentNoSkip)
    def setUp(self):
        self.tree = core.build(cwd=FILES, root=FILES)
        self.document = self.tree.find_document("REQ")
        self.temp = tempfile.mkdtemp()

    def tearDown(self):
        if os.path.exists(self.temp):
            shutil.rmtree(self.temp)

    def test_publish_html(self):
        """Verify an HTML file can be created."""
        path = os.path.join(self.temp, "published.html")
        # Act
        path2 = core.publisher.publish(self.document, path, ".html")
        # Assert
        self.assertIs(path, path2)
        filePath = os.path.join(self.temp, "documents", "published.html")
        self.assertTrue(os.path.isfile(filePath))

    def test_publish_bad_link(self):
        """Verify a tree can be published with bad links."""
        item = self.document.add_item()
        try:
            item.link("badlink")
            dirpath = os.path.join(self.temp, "html")
            # Act
            dirpath2 = core.publisher.publish(self.tree, dirpath)
            # Assert
            self.assertIs(dirpath, dirpath2)
        finally:
            item.delete()

    def test_lines_text_document(self):
        """Verify text can be published from a document."""
        path = os.path.join(FILES, "published.txt")
        expected = common.read_text(path)
        # Act
        lines = core.publisher.publish_lines(self.document, ".txt")
        text = "".join(line + "\n" for line in lines)
        # Assert
        self.assertEqual(expected, text)
        common.write_text(text, path)

    @patch("doorstop.settings.PUBLISH_CHILD_LINKS", False)
    def test_lines_text_document_without_child_links(self):
        """Verify text can be published from a document w/o child links."""
        path = os.path.join(FILES, "published2.txt")
        expected = common.read_text(path)
        # Act
        lines = core.publisher.publish_lines(self.document, ".txt")
        text = "".join(line + "\n" for line in lines)
        # Assert
        self.assertEqual(expected, text)
        common.write_text(text, path)

    def test_lines_markdown_document(self):
        """Verify Markdown can be published from a document."""
        path = os.path.join(FILES, "published.md")
        expected = common.read_text(path)
        # Act
        lines = core.publisher.publish_lines(self.document, ".md")
        text = "".join(line + "\n" for line in lines)
        # Assert
        self.assertEqual(expected, text)
        common.write_text(text, path)

    @patch("doorstop.settings.PUBLISH_CHILD_LINKS", False)
    def test_lines_markdown_document_without_child_links(self):
        """Verify Markdown can be published from a document w/o child links."""
        path = os.path.join(FILES, "published2.md")
        expected = common.read_text(path)
        # Act
        lines = core.publisher.publish_lines(self.document, ".md")
        text = "".join(line + "\n" for line in lines)
        # Assert
        self.assertEqual(expected, text)
        common.write_text(text, path)

    def test_lines_html_document_linkify(self):
        """Verify HTML can be published from a document."""
        path = os.path.join(FILES, "published.html")
        expected = common.read_text(path)
        # Act
        lines = core.publisher.publish_lines(
            self.document, ".html", linkify=True, toc=True
        )
        actual = "".join(line + "\n" for line in lines)
        # Assert
        if actual != expected:
            common.log.error(f"Published content changed: {path}")
        common.write_text(actual, path)
        self.assertEqual(expected, actual)

    @patch("doorstop.settings.PUBLISH_CHILD_LINKS", False)
    def test_lines_html_document_without_child_links(self):
        """Verify HTML can be published from a document w/o child links."""
        path = os.path.join(FILES, "published2.html")
        expected = common.read_text(path)
        # Act
        lines = core.publisher.publish_lines(self.document, ".html", toc=True)
        actual = "".join(line + "\n" for line in lines)
        # Assert
        if actual != expected:
            common.log.error(f"Published content changed: {path}")
        common.write_text(actual, path)
        self.assertEqual(expected, actual)


class TestModule(unittest.TestCase):
    """Integration tests for the doorstop.core module."""

    def setUp(self):
        """Reset the internal tree."""
        _clear_tree()

    def test_find_document(self):
        """Verify documents can be found using a convenience function."""
        # Cache miss
        document = core.find_document("req")
        self.assertIsNot(None, document)
        # Cache hit
        document2 = core.find_document("req")
        self.assertIs(document2, document)

    def test_find_item(self):
        """Verify items can be found using a convenience function."""
        # Cache miss
        item = core.find_item("req1")
        self.assertIsNot(None, item)
        # Cache hit
        item2 = core.find_item("req1")
        self.assertIs(item2, item)


# helper functions ###########################################################


def log_data(expected, actual):
    """Log list values."""
    for index, (e, a) in enumerate(zip(expected, actual)):
        logging.debug(
            "\n{i} expected:\n{e}\n{i} actual:\n{a}".format(
                i=index, e=pprint.pformat(e), a=pprint.pformat(a)
            )
        )


def read_yml(path):
    """Return a dictionary of items from a YAML file."""
    text = common.read_text(path)
    data = yaml.load(text, Loader=yaml.SafeLoader)
    return data


def read_csv(path, delimiter=","):
    """Return a list of rows from a CSV file."""
    rows = []
    try:
        with open(path, "r", newline="", encoding="utf-8") as stream:
            reader = csv.reader(stream, delimiter=delimiter)
            for row in reader:
                rows.append(row)
    except FileNotFoundError:
        logging.warning("file not found: {}".format(path))
    return rows


def read_xlsx(path):
    """Return a list of workbook data from an XLSX file."""
    data = []

    try:
        workbook = openpyxl.load_workbook(path)
    except openpyxl.exceptions.InvalidFileException:
        logging.warning("file not found: {}".format(path))
    else:
        worksheet = workbook.active
        for row in worksheet.rows:
            for cell in row:
                values = (
                    cell.value,
                    cell.style,
                    worksheet.column_dimensions[
                        openpyxl.utils.get_column_letter(cell.column)
                    ].width,
                )
                data.append(values)
        data.append(worksheet.auto_filter.ref)

    return data


def move_file(src, dst):
    """Move a file from one path to another."""
    common.delete(dst)
    shutil.move(src, dst)
