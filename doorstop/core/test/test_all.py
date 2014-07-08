"""Tests for the doorstop.core package."""

import unittest
from unittest.mock import patch, Mock

import os
import csv
import tempfile
import shutil
import pprint
import logging

import yaml
# TODO: openpyxl has false positives with pylint
import openpyxl  # pylint: disable=F0401

from doorstop import core
from doorstop.common import DoorstopWarning, DoorstopError
from doorstop.core.builder import _get_tree, _clear_tree

from doorstop.core.test import ENV, REASON, ROOT, FILES, EMPTY, SYS
from doorstop.core.test import DocumentNoSkip

# Whenever the export format is changed:
#  1. set CHECK_EXPORTED_CONTENT to False
#  2. re-run all tests
#  3. manually verify the newly exported content is correct
#  4. set CHECK_EXPORTED_CONTENT to True
CHECK_EXPORTED_CONTENT = True

# Whenever the publish format is changed:
#  1. set CHECK_PUBLISHED_CONTENT to False
#  2. re-run all tests
#  3. manually verify the newly published content is correct
#  4. set CHECK_PUBLISHED_CONTENT to True
CHECK_PUBLISHED_CONTENT = True


@unittest.skipUnless(os.getenv(ENV), REASON)  # pylint: disable=R0904
class TestItem(unittest.TestCase):  # pylint: disable=R0904

    """Integration tests for the Item class."""  # pylint: disable=C0103

    def setUp(self):
        self.path = os.path.join(FILES, 'REQ001.yml')
        with open(self.path, 'r') as item:
            self.backup = item.read()
        self.item = core.Item(self.path)

    def tearDown(self):
        with open(self.path, 'w') as item:
            item.write(self.backup)

    def test_save_load(self):
        """Verify an item can be saved and loaded from a file."""
        self.item.level = '1.2.3'
        self.item.text = "Hello, world!"
        self.item.links = ['SYS001', 'SYS002']
        item2 = core.Item(os.path.join(FILES, 'REQ001.yml'))
        self.assertEqual((1, 2, 3), item2.level)
        self.assertEqual("Hello, world!", item2.text)
        self.assertEqual(['SYS001', 'SYS002'], item2.links)

    def test_find_ref(self):
        """Verify an item's external reference can be found."""

        def skip(path):
            """Skip exported content."""
            return path.endswith(".csv") or path.endswith(".tsv")

        item = core.Item(os.path.join(FILES, 'REQ003.yml'))
        path, line = item.find_ref(skip=skip)
        relpath = os.path.relpath(os.path.join(FILES, 'external', 'text.txt'),
                                  ROOT)
        self.assertEqual(relpath, path)
        self.assertEqual(3, line)

    def test_find_ref_error(self):
        """Verify an error occurs when no external reference found."""
        self.item.ref = "not found".replace(' ', '')  # avoids self match
        self.assertRaises(DoorstopError, self.item.find_ref)


@unittest.skipUnless(os.getenv(ENV), REASON)  # pylint: disable=R0904
class TestDocument(unittest.TestCase):  # pylint: disable=R0904

    """Integration tests for the Document class."""  # pylint: disable=C0103

    def setUp(self):
        self.document = core.Document(FILES, root=ROOT)

    def tearDown(self):
        """Clean up temporary files."""
        for filename in os.listdir(EMPTY):
            path = os.path.join(EMPTY, filename)
            os.remove(path)

    def test_load(self):
        """Verify a document can be loaded from a directory."""
        doc = core.Document(FILES)
        self.assertEqual('REQ', doc.prefix)
        self.assertEqual(2, doc.digits)
        self.assertEqual(5, len(doc.items))

    def test_new(self):
        """Verify a new document can be created."""
        document = core.Document.new(None,
                                     EMPTY, FILES,
                                     prefix='SYS', digits=4)
        self.assertEqual('SYS', document.prefix)
        self.assertEqual(4, document.digits)
        self.assertEqual(0, len(document.items))

    @patch('doorstop.settings.REORDER', False)
    def test_validate(self):
        """Verify a document can be validated."""
        self.assertTrue(self.document.validate())

    @patch('doorstop.settings.REORDER', False)
    def test_issues_count(self):
        """Verify a number of issues are found in a document."""
        issues = self.document.issues
        for issue in self.document.issues:
            logging.info(repr(issue))
        self.assertEqual(8, len(issues))

    @patch('doorstop.settings.REORDER', False)
    def test_issues_duplicate_level(self):
        """Verify duplicate item levels are detected."""
        expect = DoorstopWarning("duplicate level: 2.1 (REQ002, REQ2-001)")
        for issue in self.document.issues:
            logging.info(repr(issue))
            if type(issue) == type(expect) and issue.args == expect.args:
                break
        else:
            self.fail("issue not found: {}".format(expect))

    @patch('doorstop.settings.REORDER', False)
    def test_issues_skipped_level(self):
        """Verify skipped item levels are detected."""
        expect = DoorstopWarning("skipped level: 1.4 (REQ003), 1.6 (REQ004)")
        for issue in self.document.issues:
            logging.info(repr(issue))
            if type(issue) == type(expect) and issue.args == expect.args:
                break
        else:
            self.fail("issue not found: {}".format(expect))

    def test_add_item_with_reordering(self):
        """Verify an item can be inserted into a document."""
        document = core.Document.new(None,
                                     EMPTY, FILES,
                                     prefix='TMP')
        item_1_0 = document.add_item()
        item_3_0 = document.add_item()  # will get displaced
        item_2_0 = document.add_item(level='2.0')
        self.assertEqual((1, 0), item_1_0.level)
        self.assertEqual((2, 0), item_2_0.level)
        self.assertEqual((3, 0), item_3_0.level)

    def test_remove_item_with_reordering(self):
        """Verify an item can be removed fraom a document."""
        document = core.Document.new(None,
                                     EMPTY, FILES,
                                     prefix='TMP')
        item_1_0 = document.add_item()
        item_3_0 = document.add_item()  # to be removed
        item_2_0 = document.add_item()  # will get relocated
        document.remove_item(item_3_0)
        self.assertEqual((1, 0), item_1_0.level)
        self.assertEqual((2, 0), item_2_0.level)

    def test_reorder(self):
        """Verify a document's order can be corrected."""
        document = core.Document.new(None,
                                     EMPTY, FILES,
                                     prefix='TMP')
        document.add_item(level='2.0', reorder=False)
        document.add_item(level='2.1', reorder=False)
        document.add_item(level='2.1', reorder=False)
        document.add_item(level='2.5', reorder=False)
        document.add_item(level='4.5', reorder=False)
        document.add_item(level='4.7', reorder=False)
        document.reorder()
        expected = [(2, 0), (2, 1), (2, 2), (2, 3), (3, 1), (3, 2)]
        actual = [item.level for item in document.items]
        self.assertListEqual(expected, actual)

    def test_reorder_with_keep(self):
        """Verify a document's order can be corrected with a kept level."""
        document = core.Document.new(None,
                                     EMPTY, FILES,
                                     prefix='TMP')
        document.add_item(level='1.0', reorder=False)
        item = document.add_item(level='1.0', reorder=False)
        document.add_item(level='1.0', reorder=False)
        document.reorder(keep=item)
        expected = [(1, 0), (2, 0), (3, 0)]
        actual = [item.level for item in document.items]
        self.assertListEqual(expected, actual)
        self.assertEqual((1, 0), item.level)

    def test_reorder_with_start(self):
        """Verify a document's order can be corrected with a given start."""
        document = core.Document.new(None,
                                     EMPTY, FILES,
                                     prefix='TMP')
        document.add_item(level='2.0', reorder=False)
        document.add_item(level='2.1', reorder=False)
        document.add_item(level='2.1', reorder=False)
        document.add_item(level='2.5', reorder=False)
        document.add_item(level='4.0', reorder=False)
        document.add_item(level='4.7', reorder=False)
        document.reorder(start=(1, 0))
        expected = [(1, 0), (1, 1), (1, 2), (1, 3), (2, 0), (2, 1)]
        actual = [item.level for item in document.items]
        self.assertListEqual(expected, actual)

    @patch('doorstop.settings.REORDER', True)
    def test_validate_with_reordering(self):
        """Verify a document's order is corrected during validation."""
        document = core.Document.new(None,
                                     EMPTY, FILES,
                                     prefix='TMP')
        document.add_item(level='1.0', reorder=False)
        document.add_item(level='1.1', reorder=False)
        document.add_item(level='1.2.0', reorder=False)
        document.add_item(level='1.2.5', reorder=False)
        document.add_item(level='3.2.1', reorder=False)
        document.add_item(level='3.3', reorder=False)
        self.assertTrue(document.validate())
        expected = [(1, 0), (1, 1), (1, 2, 0), (1, 2, 1), (2, 1, 1), (2, 2)]
        actual = [item.level for item in document.items]
        self.assertListEqual(expected, actual)


@unittest.skipUnless(os.getenv(ENV), REASON)  # pylint: disable=R0904
class TestTree(unittest.TestCase):  # pylint: disable=R0904

    """Integration tests for the core.Tree class."""

    def setUp(self):
        self.path = os.path.join(FILES, 'REQ001.yml')
        with open(self.path, 'r') as item:
            self.backup = item.read()
        self.item = core.Item(self.path)
        self.tree = core.Tree(core.Document(SYS))
        self.tree._place(core.Document(FILES))  # pylint: disable=W0212

    def tearDown(self):
        with open(self.path, 'w') as item:
            item.write(self.backup)

    @patch('doorstop.settings.REORDER', False)
    @patch('doorstop.core.document.Document', DocumentNoSkip)
    def test_validate_invalid_link(self):
        """Verify a tree is invalid with a bad link."""
        self.item.link('SYS003')
        tree = core.build(FILES, root=FILES)
        self.assertIsInstance(tree, core.Tree)
        self.assertFalse(tree.validate())

    @patch('doorstop.settings.REORDER', False)
    def test_validate_long(self):
        """Verify trees can be checked."""
        logging.info("tree: {}".format(self.tree))
        self.assertTrue(self.tree.validate())


@unittest.skipUnless(os.getenv(ENV), REASON)  # pylint: disable=R0904
class TestEditor(unittest.TestCase):  # pylint: disable=R0904

    """Integrations tests for the editor module."""  # pylint: disable=C0103


@unittest.skipUnless(os.getenv(ENV), REASON)  # pylint: disable=R0902,R0904
class TestImporter(unittest.TestCase):  # pylint: disable=R0904

    """Integrations tests for the importer module."""  # pylint: disable=C0103

    def setUp(self):
        # Create a temporary mock working copy
        self.cwd = os.getcwd()
        self.temp = tempfile.mkdtemp()
        os.chdir(self.temp)
        open(".mockvcs", 'w').close()
        # Create default document attributes
        self.prefix = 'PREFIX'
        self.root = self.temp
        self.path = os.path.join(self.root, 'DIRECTORY')
        self.parent = 'PARENT_PREFIX'
        # Create default item attributes
        self.identifier = 'PREFIX-00042'
        # Load an actual document
        self.document = core.Document(FILES, root=ROOT)
        # Ensure the tree is reloaded
        _clear_tree()

    def tearDown(self):
        os.chdir(self.cwd)
        shutil.rmtree(self.temp)

    def test_import_yml(self):
        """Verify items can be imported from a YAML file."""
        path = os.path.join(self.temp, 'exported.yml')
        core.exporter.export(self.document, path)
        _path = os.path.join(self.temp, 'imports', 'req')
        _tree = _get_tree()
        document = _tree.create_document(_path, 'REQ')
        # Act
        core.importer.import_file(path, document)
        # Assert
        expected = [item.data for item in self.document.items]
        actual = [item.data for item in document.items]
        log_data(expected, actual)
        self.assertListEqual(expected, actual)

    def test_import_csv(self):
        """Verify items can be imported from a CSV file."""
        path = os.path.join(self.temp, 'exported.csv')
        core.exporter.export(self.document, path)
        _path = os.path.join(self.temp, 'imports', 'req')
        _tree = _get_tree()
        document = _tree.create_document(_path, 'REQ')
        # Act
        core.importer.import_file(path, document)
        # Assert
        expected = [item.data for item in self.document.items]
        actual = [item.data for item in document.items]
        log_data(expected, actual)
        self.assertListEqual(expected, actual)

    def test_import_tsv(self):
        """Verify items can be imported from a TSV file."""
        path = os.path.join(self.temp, 'exported.tsv')
        core.exporter.export(self.document, path)
        _path = os.path.join(self.temp, 'imports', 'req')
        _tree = _get_tree()
        document = _tree.create_document(_path, 'REQ')
        # Act
        core.importer.import_file(path, document)
        # Assert
        expected = [item.data for item in self.document.items]
        actual = [item.data for item in document.items]
        log_data(expected, actual)
        self.assertListEqual(expected, actual)

    def test_import_xlsx(self):
        """Verify items can be imported from an XLSX file."""
        path = os.path.join(self.temp, 'exported.xlsx')
        core.exporter.export(self.document, path)
        _path = os.path.join(self.temp, 'imports', 'req')
        _tree = _get_tree()
        document = _tree.create_document(_path, 'REQ')
        # Act
        core.importer.import_file(path, document)
        # Assert
        expected = [item.data for item in self.document.items]
        actual = [item.data for item in document.items]
        log_data(expected, actual)
        self.assertListEqual(expected, actual)

    def test_create_document(self):
        """Verify a new document can be created to import items."""
        document = core.importer.create_document(self.prefix, self.path)
        self.assertEqual(self.prefix, document.prefix)
        self.assertEqual(self.path, document.path)

    def test_create_document_with_unknown_parent(self):
        """Verify a new document can be created with an unknown parent."""
        # Verify the document does not already exist
        self.assertRaises(DoorstopError, core.find_document, self.prefix)
        # Import a document
        document = core.importer.create_document(self.prefix, self.path,
                                                 parent=self.parent)
        # Verify the imported document's attributes are correct
        self.assertEqual(self.prefix, document.prefix)
        self.assertEqual(self.path, document.path)
        self.assertEqual(self.parent, document.parent)
        # Verify the imported document can be found
        document2 = core.find_document(self.prefix)
        self.assertIs(document, document2)

    def test_create_document_already_exists(self):
        """Verify non-parent exceptions are re-raised."""
        # Create a document
        core.importer.create_document(self.prefix, self.path)
        # Attempt to create the same document
        self.assertRaises(DoorstopError, core.importer.create_document,
                          self.prefix, self.path)

    def test_add_item(self):
        """Verify an item can be imported into a document."""
        # Create a document
        core.importer.create_document(self.prefix, self.path)
        # Verify the item does not already exist
        self.assertRaises(DoorstopError, core.find_item, self.identifier)
        # Import an item
        item = core.importer.add_item(self.prefix, self.identifier)
        # Verify the item's attributes are correct
        self.assertEqual(self.identifier, item.id)
        # Verify the item can be found
        item2 = core.find_item(self.identifier)
        self.assertIs(item, item2)
        # Verify the item is contained in the document
        document = core.find_document(self.prefix)
        self.assertIn(item, document.items)

    def test_add_item_with_attrs(self):
        """Verify an item with attributes can be imported into a document."""
        # Create a document
        core.importer.create_document(self.prefix, self.path)
        # Import an item
        attrs = {'text': "Item text", 'ext1': "Extended 1"}
        item = core.importer.add_item(self.prefix, self.identifier,
                                      attrs=attrs)
        # Verify the item is correct
        self.assertEqual(self.identifier, item.id)
        self.assertEqual(attrs['text'], item.text)
        self.assertEqual(attrs['ext1'], item.get('ext1'))


@unittest.skipUnless(os.getenv(ENV) or not CHECK_EXPORTED_CONTENT, REASON)  # pylint: disable=R0904
class TestExporter(unittest.TestCase):  # pylint: disable=R0904

    """Integration tests for the doorstop.core.exporter module."""  # pylint: disable=C0103

    maxDiff = None

    def setUp(self):
        self.document = core.Document(FILES, root=ROOT)
        self.temp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.temp)

    def test_export_yml(self):
        """Verify a document can be exported as a YAML file."""
        path = os.path.join(FILES, 'exported.yml')
        temp = os.path.join(self.temp, 'exported.yml')
        expected = read_yml(path)
        # Act
        path2 = core.exporter.export(self.document, temp)
        # Assert
        self.assertIs(temp, path2)
        if CHECK_EXPORTED_CONTENT:
            actual = read_yml(temp)
            self.assertEqual(expected, actual)
        move_file(temp, path)

    def test_export_csv(self):
        """Verify a document can be exported as a CSV file."""
        path = os.path.join(FILES, 'exported.csv')
        temp = os.path.join(self.temp, 'exported.csv')
        expected = read_csv(path)
        # Act
        path2 = core.exporter.export(self.document, temp)
        # Assert
        self.assertIs(temp, path2)
        if CHECK_EXPORTED_CONTENT:
            actual = read_csv(temp)
            self.assertEqual(expected, actual)
        move_file(temp, path)

    def test_export_tsv(self):
        """Verify a document can be exported as a TSV file."""
        path = os.path.join(FILES, 'exported.tsv')
        temp = os.path.join(self.temp, 'exported.tsv')
        expected = read_csv(path, delimiter='\t')
        # Act
        path2 = core.exporter.export(self.document, temp)
        # Assert
        self.assertIs(temp, path2)
        if CHECK_EXPORTED_CONTENT:
            actual = read_csv(temp, delimiter='\t')
            self.assertEqual(expected, actual)
        move_file(temp, path)

    def test_export_xlsx(self):
        """Verify a document can be exported as an XLSX file."""
        path = os.path.join(FILES, 'exported.xlsx')
        temp = os.path.join(self.temp, 'exported.xlsx')
        expected = read_xlsx(path)
        # Act
        path2 = core.exporter.export(self.document, temp)
        # Assert
        self.assertIs(temp, path2)
        if CHECK_EXPORTED_CONTENT:
            actual = read_xlsx(temp)
            self.assertEqual(expected, actual)
        else:  # binary file always changes, only copy when not checking
            move_file(temp, path)


@unittest.skipUnless(os.getenv(ENV) or not CHECK_PUBLISHED_CONTENT, REASON)  # pylint: disable=R0904
class TestPublisher(unittest.TestCase):  # pylint: disable=R0904

    """Integration tests for the doorstop.core.publisher module."""  # pylint: disable=C0103

    maxDiff = None

    @patch('doorstop.core.document.Document', DocumentNoSkip)
    def setUp(self):
        self.tree = core.build(cwd=FILES, root=FILES)
        # self.document = core.Document(FILES, root=ROOT)
        self.document = self.tree.find_document('REQ')
        self.temp = tempfile.mkdtemp()

    def tearDown(self):
        if os.path.exists(self.temp):
            shutil.rmtree(self.temp)

    def test_publish_html(self):
        """Verify an HTML file can be created."""
        path = os.path.join(self.temp, 'published.html')
        # Act
        path2 = core.publisher.publish(self.document, path, '.html')
        # Assert
        self.assertIs(path, path2)
        self.assertTrue(os.path.isfile(path))

    def test_publish_empty_tree(self):
        """Verify a directory is not created when no documents to publish"""
        mock_tree = Mock()
        mock_tree.documents = []
        dirpath = os.path.join(self.temp, 'gen')
        # Act
        path2 = core.publisher.publish(mock_tree, dirpath, index=True)
        # Assert
        self.assertIs(None, path2)
        self.assertFalse(os.path.exists(dirpath))

    def test_publish_bad_link(self):
        """Verify a tree can be published with bad links."""
        item = self.document.add_item()
        try:
            item.link('badlink')
            dirpath = os.path.join(self.temp, 'html')
            # Act
            dirpath2 = core.publisher.publish(self.tree, dirpath)
            # Assert
            self.assertIs(dirpath, dirpath2)
        finally:
            item.delete()

    def test_lines_text_document(self):
        """Verify text can be published from a document."""
        path = os.path.join(FILES, 'published.txt')
        expected = open(path, encoding='utf-8').read()
        # Act
        lines = core.publisher.publish_lines(self.document, '.txt')
        text = ''.join(line + '\n' for line in lines)
        # Assert
        if CHECK_PUBLISHED_CONTENT:
            self.assertEqual(expected, text)
        with open(path, 'w', encoding='utf-8') as outfile:
            outfile.write(text)

    @patch('doorstop.settings.PUBLISH_CHILD_LINKS', False)
    def test_lines_text_document_without_child_links(self):
        """Verify text can be published from a document w/o child links."""
        path = os.path.join(FILES, 'published2.txt')
        expected = open(path, encoding='utf-8').read()
        # Act
        lines = core.publisher.publish_lines(self.document, '.txt')
        text = ''.join(line + '\n' for line in lines)
        # Assert
        if CHECK_PUBLISHED_CONTENT:
            self.assertEqual(expected, text)
        with open(path, 'w', encoding='utf-8') as outfile:
            outfile.write(text)

    def test_lines_markdown_document(self):
        """Verify Markdown can be published from a document."""
        path = os.path.join(FILES, 'published.md')
        expected = open(path, encoding='utf-8').read()
        # Act
        lines = core.publisher.publish_lines(self.document, '.md')
        text = ''.join(line + '\n' for line in lines)
        # Assert
        if CHECK_PUBLISHED_CONTENT:
            self.assertEqual(expected, text)
        with open(path, 'w', encoding='utf-8') as outfile:
            outfile.write(text)

    @patch('doorstop.settings.PUBLISH_CHILD_LINKS', False)
    def test_lines_markdown_document_without_child_links(self):
        """Verify Markdown can be published from a document w/o child links."""
        path = os.path.join(FILES, 'published2.md')
        expected = open(path, encoding='utf-8').read()
        # Act
        lines = core.publisher.publish_lines(self.document, '.md')
        text = ''.join(line + '\n' for line in lines)
        # Assert
        if CHECK_PUBLISHED_CONTENT:
            self.assertEqual(expected, text)
        with open(path, 'w', encoding='utf-8') as outfile:
            outfile.write(text)

    def test_lines_html_document_linkify(self):
        """Verify HTML can be published from a document."""
        path = os.path.join(FILES, 'published.html')
        expected = open(path, encoding='utf-8').read()
        # Act
        lines = core.publisher.publish_lines(self.document, '.html',
                                             linkify=True)
        text = ''.join(line + '\n' for line in lines)
        # Assert
        if CHECK_PUBLISHED_CONTENT:
            self.assertEqual(expected, text)
        with open(path, 'w', encoding='utf-8') as outfile:
            outfile.write(text)

    @patch('doorstop.settings.PUBLISH_CHILD_LINKS', False)
    def test_lines_html_document_without_child_links(self):
        """Verify HTML can be published from a document w/o child links."""
        path = os.path.join(FILES, 'published2.html')
        expected = open(path, encoding='utf-8').read()
        # Act
        lines = core.publisher.publish_lines(self.document, '.html')
        text = ''.join(line + '\n' for line in lines)
        # Assert
        if CHECK_PUBLISHED_CONTENT:
            self.assertEqual(expected, text)
        with open(path, 'w', encoding='utf-8') as outfile:
            outfile.write(text)


@unittest.skipUnless(os.getenv(ENV), REASON)  # pylint: disable=R0904
class TestModule(unittest.TestCase):  # pylint: disable=R0904

    """Integration tests for the doorstop.core module."""  # pylint: disable=C0103

    def setUp(self):
        """Reset the internal tree."""
        _clear_tree()

    def test_find_document(self):
        """Verify documents can be found using a convenience function."""
        # Cache miss
        document = core.find_document('req')
        self.assertIsNot(None, document)
        # Cache hit
        document2 = core.find_document('req')
        self.assertIs(document2, document)

    def test_find_item(self):
        """Verify items can be found using a convenience function."""
        # Cache miss
        item = core.find_item('req1')
        self.assertIsNot(None, item)
        # Cache hit
        item2 = core.find_item('req1')
        self.assertIs(item2, item)


# helper functions ###########################################################


def log_data(expected, actual):
    """Log list values."""
    for index, (evalue, avalue) in enumerate(zip(expected, actual)):
        logging.debug("\n{i} expected:\n{e}\n{i} actual:\n{a}".format(
            i=index,
            e=pprint.pformat(evalue),
            a=pprint.pformat(avalue)))


def read_yml(path):
    """Return a dictionary of items from a YAML file."""
    with open(path, 'r', encoding='utf-8') as stream:
        text = stream.read()
    data = yaml.load(text)
    return data


def read_csv(path, delimiter=','):
    """Return a list of rows from a CSV file."""
    rows = []
    try:
        with open(path, 'r', newline='', encoding='utf-8') as stream:
            reader = csv.reader(stream, delimiter=delimiter)
            for row in reader:
                rows.append(row)
    except FileNotFoundError:
        logging.warning("file not found: {}".format(path))
    return rows


def read_xlsx(path):
    """Return a list of workbook data from an XLSX file."""
    data = []

    try:
        workbook = openpyxl.load_workbook(path)
    except openpyxl.exceptions.InvalidFileException:
        logging.warning("file not found: {}".format(path))
    else:
        worksheet = workbook.active
        for row in worksheet.rows:
            for cell in row:
                values = (cell.value,
                          cell.style,
                          worksheet.column_dimensions[cell.column].width)
                data.append(values)
        data.append(worksheet.auto_filter.ref)

    return data


def move_file(src, dst):
    """Move a file from one path to another."""
    try:
        os.remove(dst)
    except FileNotFoundError:
        pass
    shutil.move(src, dst)
